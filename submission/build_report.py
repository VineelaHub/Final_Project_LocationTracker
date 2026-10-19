from pathlib import Path
import shutil, re, zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent.parent
out = root / 'submission'
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.8)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'; normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for name in ['Title','Subtitle','Heading 1','Heading 2']:
    doc.styles[name].font.color.rgb = RGBColor(0,0,0)
doc.styles['Title'].font.size = Pt(25)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 2'].font.size = Pt(12)
def p(text): doc.add_paragraph(text)
def h(text): doc.add_heading(text,1)
def page(): doc.add_page_break()
doc.add_paragraph('Location Tracker Application', 'Title')
doc.add_paragraph('Final Software Development Project', 'Subtitle')
p('Student name: [Enter your name]\nCourse: Multiplatform App Development\nInstructor: [Enter instructor name]\nSubmission date: [Enter date]')
h('Project overview')
p('This project implements a location tracking application in C# and .NET MAUI. It captures GPS observations, stores them in a local SQLite database, and extends the .NET MAUI Maps component with translucent circle overlays to visualize location density. The Android application uses Google Maps as its native map provider.')
p('The interface follows the supplied reference with a purple MainView title bar and a map occupying the rest of the page. A small New route control allows a fresh recording without restoring the earlier large button panel. Tracking begins automatically after the page opens and location permission is granted.')
h('Assignment requirements')
for s in ['C# and .NET MAUI source code with comments describing the main responsibilities.', 'GPS location collection and local SQLite persistence.', 'A heat map implemented on the .NET MAUI Maps component.', 'Android manifest with Internet and location permissions and an API key configuration entry.', 'Application output screenshots in this report and a GitHub repository URL for the source code.']:
    doc.add_paragraph(s, 'List Bullet')
h('Repository')
p('GitHub repository URL: [PASTE YOUR GITHUB REPOSITORY URL HERE]')
p('The accompanying Source folder contains the project and platform configuration. The distributable manifest uses an API key placeholder; a valid restricted Maps SDK for Android key must be configured locally before running the map.')
page()
h('Application implementation')
doc.add_heading('Location collection',2)
p('MainPage requests LocationWhenInUse permission and uses Geolocation.GetLocationAsync with medium accuracy and a ten second timeout. While tracking is active, the loop requests a position approximately every fifteen seconds after the previous capture completes. Automatic observations are saved when the device has moved at least ten metres from the last saved point. This reduces repeated observations while stationary.')
p('Tracking is canceled when the page disappears. The implementation is intended for foreground use and does not provide an Android background location service. Operating system pause and resume behavior should be checked on the target device.')
doc.add_heading('SQLite storage',2)
p('LocationPoint stores an automatically assigned identifier, latitude, longitude, optional accuracy in metres, and a UTC timestamp. LocationDatabase initializes the table once under a semaphore and exposes asynchronous save, read, and clear methods. Observations are retrieved in timestamp order.')
p('MauiProgram registers the database service with a locations.db3 file in FileSystem.AppDataDirectory. Saved history is loaded when the page opens, so closing the app does not intentionally erase the database. Persistence across a full restart should be verified during the final demonstration.')
doc.add_heading('Heat map rendering',2)
p('HeatMapRenderer clears the existing map overlays and draws three circles at each saved observation. The outer orange halo has a radius of 65 metres, the middle red orange circle 42 metres, and the dark red center 25 metres. Translucent fills allow nearby observations to overlap and increase visual intensity. All outer halos are drawn before the centers so they do not obscure neighboring centers.')
p('This is a circle based approximation of a density heat map. It is not a continuous raster heat map or a road routing engine. Closely spaced samples create a red trail; larger gaps remain visible when the device moves farther between captures.')
doc.add_heading('Starting a new route',2)
p('Changing the simulated route in the emulator does not tell the application to delete SQLite history. The New route control asks for confirmation, cancels tracking, waits for any active tracking task to finish, clears saved observations, resets the last position, redraws the map, and restarts tracking. Waiting for the task prevents a previous capture from restoring an old point after the reset.')
page()
h('Build and demonstration')
p('The project targets .NET 9 Android and iOS and references Microsoft.Maui.Controls 9.0.120, Microsoft.Maui.Controls.Maps, and sqlite-net-pcl 1.9.172. Android was built and demonstrated in the emulator. iOS is included as a target but has not been built or tested in this submission workflow; it requires a compatible Mac build environment.')
doc.add_heading('Observed results and remaining checks',2)
for s in ['Android build and run: successful build output was produced. The earlier full build reported SQLite native library page size warnings.', 'Map display and collection: emulator screenshots showed a working Google map and a route with 23 saved points.', 'Reference layout and heat map: subsequent screenshots showed the purple MainView bar, full map, and red orange observation overlays.', 'New route reset: the source was updated to clear history safely and passed compilation. Verify that the previous trail disappears after confirming New route.', 'Persistence: close and reopen the app and confirm the saved trail reloads. Record the actual result in the final screenshot caption if demonstrated.']:
    doc.add_paragraph(s,'List Bullet')
doc.add_heading('Reproducing the demonstration',2)
for s in ['Install the .NET 9 SDK, a compatible MAUI Android workload, a Java development kit, and Android SDK platform 35. Start an Android emulator with Google Play services.', 'Enable Maps SDK for Android in Google Cloud. Put a restricted key in Platforms/Android/AndroidManifest.xml. Match package edu.cumberlands.locationtracker and the SHA 1 of the actual signing certificate.', 'Restore and build the Android target using the commands in the source README. Grant location permission while using the app.', 'In the emulator Location controls, set a starting point. Confirm New route in the app, then play a simulated route. Keep the map visible while observations accumulate.', 'Capture screenshots of the completed heat map and the new route behavior. Insert them on the next page. Add the repository URL and personal details before submission.']:
    doc.add_paragraph(s,'List Number')
doc.add_heading('References',2)
for s in ['Microsoft Learn. Map control in .NET MAUI. https://learn.microsoft.com/dotnet/maui/user-interface/controls/map', 'Microsoft Learn. Geolocation in .NET MAUI. https://learn.microsoft.com/dotnet/maui/platform-integration/device/geolocation', 'Microsoft Learn. Permissions in .NET MAUI. https://learn.microsoft.com/dotnet/maui/platform-integration/appmodel/permissions', 'sqlite-net project. https://github.com/praeclarum/sqlite-net']:
    para=doc.add_paragraph(s); para.paragraph_format.space_after=Pt(4)
    for run in para.runs: run.font.size=Pt(9)
page()
h('Application output screenshots')
p('Replace the placeholders below with your own screenshots. Use the final version with MainView, the red heat map, and the New route control.')
for title, caption in [('Screenshot 1 Completed route heat map','Figure 1. Android emulator showing saved GPS observations rendered as a red orange heat trail using native .NET MAUI Maps circles.'), ('Screenshot 2 Fresh route after reset','Figure 2. New route demonstration after confirming removal of previous history and recording locations on a different simulated route.')]:
    doc.add_heading(title,2)
    table=doc.add_table(rows=1, cols=1)
    table.rows[0].height=Inches(2.4)
    table.rows[0].cells[0].text='INSERT YOUR APPLICATION SCREENSHOT HERE'
    cell=table.cell(0,0)
    cell.vertical_alignment=1
    cell.paragraphs[0].alignment=1
    shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'F2F2F2'); cell._tc.get_or_add_tcPr().append(shd)
    doc.add_paragraph(caption,'Caption')
doc.core_properties.title='Location Tracker Application'
doc.core_properties.subject='Final Software Development Project'
doc.core_properties.author=''
doc.save(out / 'Location_Tracker_Report.docx')

source=out/'Source'
for f in (root/'LocationTracker').rglob('*'):
    if not f.is_file() or any(part in {'bin','obj','.vs','.git'} for part in f.relative_to(root/'LocationTracker').parts): continue
    dest=source/'LocationTracker'/f.relative_to(root/'LocationTracker')
    dest.parent.mkdir(parents=True,exist_ok=True)
    if f.name=='AndroidManifest.xml':
        text=f.read_text(encoding='utf-8-sig')
        text=re.sub(r'(android:name="com.google.android.geo.API_KEY"\s+android:value=")[^"]*',r'\g<1>YOUR_GOOGLE_MAPS_API_KEY',text)
        dest.write_text(text,encoding='utf-8')
    else: shutil.copy2(f,dest)
readme='''# Location Tracker

C# .NET MAUI assignment app using native Maps circle overlays, foreground GPS, and SQLite.

## Interface and behavior
Purple MainView title bar, full-page map, automatic foreground tracking, and New route confirmation. Captures approximately every 15 seconds after each request completes; automatic saves require at least 10 metres of movement. Red/orange layered circle overlays approximate location density. New route deletes saved history and restarts collection.

## Android setup
Install .NET 9 SDK, compatible MAUI Android tooling, a JDK, and Android SDK platform 35. Use an emulator with Google Play services. Replace YOUR_GOOGLE_MAPS_API_KEY in LocationTracker/Platforms/Android/AndroidManifest.xml with a restricted Maps SDK for Android key. Package: edu.cumberlands.locationtracker. Configure your own signing certificate SHA-1 in Google Cloud.

From the LocationTracker folder in PowerShell:

```powershell
dotnet restore .\\LocationTracker.csproj -p:TargetFrameworks=net9.0-android
dotnet build .\\LocationTracker.csproj -f net9.0-android -p:TargetFrameworks=net9.0-android "-p:AndroidSdkDirectory=$env:LOCALAPPDATA\\Android\\Sdk" -t:Run
```

Grant location permission while using the app. In emulator Location controls, select a start, tap New route in the app, confirm, and play a route. A different emulator route alone does not clear stored app history.

## Source organization
- MainPage.cs: map layout, permission requests, capture loop, and route reset.
- Models/LocationPoint.cs: SQLite observation model.
- Services/LocationDatabase.cs: asynchronous persistence.
- Services/HeatMapRenderer.cs: layered circle heat map.
- MauiProgram.cs: Maps setup and database dependency injection.
- Platforms/Android/AndroidManifest.xml: permissions and API key placeholder.
- Platforms/iOS: iOS entry points and location usage description.

## Validation and limits
Android build and emulator screenshots verified map display and a recorded trail. The final reset change passed compilation; verify reset and persistence in the final demo. iOS is targeted but untested and requires a Mac. Tracking is not a background service. SQLite's existing native dependency produced 16 KB page-size warnings; newer-device compatibility needs separate validation. Heat rendering uses geographic circles, not a continuous raster density algorithm.

Never commit signing keystores or unrestricted API keys. The submission copy intentionally contains a key placeholder.
'''
(source/'README.md').write_text(readme,encoding='utf-8')
(source/'.gitignore').write_text('**/bin/\n**/obj/\n.vs/\n.vscode/\n*.user\n*.keystore\n*.jks\n*.apk\n*.aab\n',encoding='utf-8')
(out/'SUBMISSION_CHECKLIST.txt').write_text('''Before submission
1. Enter your name, instructor, date, and GitHub repository URL in the Word report.
2. Insert your final application screenshots in the report.
3. Test New route and reopening the app; update the report with your actual results if needed.
4. Create your GitHub repository and upload the contents of Source, including its README and platform manifests. The archive contains no real API key or signing certificate.
5. Submit the Word report and provide the GitHub repository URL. Include the source ZIP if your instructor requests a file upload.
6. Follow any course policy on acknowledging assistance and ensure you can explain your implementation.
''',encoding='utf-8')
with zipfile.ZipFile(out/'Location_Tracker_Source.zip','w',zipfile.ZIP_DEFLATED) as z:
    for f in source.rglob('*'):
        if f.is_file(): z.write(f, f.relative_to(source))
print('Created report and source archive')
